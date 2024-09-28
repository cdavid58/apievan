from datetime import datetime, timedelta
from inventory.models import Product
from company.models import Branch
from invoice.models import *
from django.db import models
import pandas as pd, os
from django.http import HttpResponse
from openpyxl import Workbook

class AccountingEntry(models.Model):
    invoice = models.ForeignKey(Invoice, on_delete=models.CASCADE)
    sequence = models.IntegerField()  # Secuencia
    date_created = models.DateField()  # Fecha de elaboración
    accounting_code = models.CharField(max_length=30)  # Código contable
    accounting_account = models.CharField(max_length=100)  # Cuenta contable
    identification = models.CharField(max_length=20)  # Identificación (Cliente)
    description = models.TextField()  # Descripción
    debit = models.FloatField(default=0)  # Débito
    credit = models.FloatField(default=0)  # Crédito
    data = []
    total = 0
    number = None
    date = None

    def __str__(self):
        return f"Entry {self.sequence} for Invoice {self.invoice.number}"

    class Meta:
        verbose_name = 'Accounting Entry'
        verbose_name_plural = 'Accounting Entries'


    @classmethod
    def create_accountingentry(cls, data, invoice):
        result = False
        message = None
        try:
            print(len(data),'Len')
            for i in range(len(data)):
                accountingentry = cls(
                    invoice = invoice,
                    sequence = data[i]['sequence'],
                    date_created = data[i]['date_created'],
                    accounting_code = data[i]['accounting_code'],
                    accounting_account = data[i]['accounting_account'],
                    identification = data[i]['identification'],
                    description = data[i]['description'],
                    debit = data[i]['debit'],
                    credit = data[i]['credit']
                )
                print(data[i])
                accountingentry.save()
                result = True
        except Exception as e:
            print(e)
            message = str(e)
        return {'result': result, 'message': message}

    @staticmethod
    def Name_Category_Data(cls):
        cls.accounting_code = None
        cls.accounting_name = None

    @staticmethod
    def Taxes(cls):
        cls.tax_19_general = 0
        cls.tax_5_general = 0
        cls.ipo_0_general = 0
        cls.tax_19_gaseosa = 0
        cls.tax_19_licor = 0
        cls.ipo_19_licor = 0
        cls.tax_19_cerveza = 0
        cls.ipo_19_cerveza = 0
        cls.tax_19_vino = 0
        cls.ipo_19_vino = 0
        cls.tax_19_cigarrillo = 0
        cls.ipo_19_cigarrillo = 0
        cls.tax_19_bolsa = 0

    @staticmethod
    def Create_Document(cls, name_doc):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        encabezado = ["Secuencia", "Fecha elaboración", "Código contable", "Cuenta contable", 
                        "Identificación", "Descripción", "Débito", "Crédito"]
        sheet.append(encabezado)
        workbook.save(f"{name_doc}.xlsx")

    @staticmethod
    def Name_Category(cls,cat, tax, number, client, base, date):

        if  tax == 19 and "GENERAL" in cat:
            cls.accounting_code = 41350101
            cls.accounting_name = 'VTA PRODU GNERAL 19%'

        elif tax == 5 and "GENERAL" in cat:
            cls.accounting_code = 41350102
            cls.accounting_name = 'VTA PROD GNERL DEL 5%'

        elif tax == 0 and "GENERAL" in cat:
            cls.accounting_code = 41350103 
            cls.accounting_name = 'VTA EXCENTA'

        elif tax == 19 and "GASEOSA" in cat:
            cls.accounting_code = 41350105
            cls.accounting_name = 'VTA GASEOSA'

        elif tax == 5 and "LICOR" in cat:
            cls.accounting_code = 41350106
            cls.accounting_name = 'VTA LICOR GNR DEL 5'

        elif tax == 19 and "CERVEZA" in cat:
            cls.accounting_code = 41350107
            cls.accounting_name = 'VTA CERVEZA DEL 19'

        elif tax == 5 and "VINO" in cat:
            cls.accounting_code = 41350108 
            cls.accounting_name = 'VTA VINO DEL 5'

        elif tax == 19 and "CIGARRILLO" in cat:
            cls.accounting_code = 41350109 
            cls.accounting_name = 'VTA CIGARRILLO  DEL19'

        cls.data.append({
            'sequence':number,
            'date_created': date,
            'accounting_code': cls.accounting_code,
            'accounting_account': cls.accounting_name,
            'identification':client,
            'description':'Producto',
            'debit':0,
            'credit': base
        })

    @staticmethod
    def Calculate_Taxes(cls,ipo,tax,value,cat):
        if 19 == tax and cat == "GENERAL":
            cls.tax_19_general += round(float(value))

        elif 5 == tax and cat == "GENERAL":
            cls.tax_5_general += round(float(value))

        elif 0 == tax and cat == "GENERAL":
            cls.ipo_0_general += round(float(ipo))

        elif 19 == tax and cat == "GASEOSA":
            cls.tax_19_gaseosa += round(float(value))

        elif 5 == tax and cat == "LICOR":
            cls.tax_19_licor += round(float(value))
            cls.ipo_19_licor += round(float(ipo))
        
        elif 19 == tax and cat == "CERVEZA":
            cls.tax_19_cerveza += round(float(value))
            cls.ipo_19_cerveza += round(float(ipo))

        elif 5 == tax and cat == "VINO":
            cls.tax_19_vino += round(float(value))
            cls.ipo_19_vino += round(float(ipo))

        elif 19 == tax and cat == "CIGARRILLO":
            cls.tax_19_cigarrillo += round(float(value))
            cls.ipo_19_cigarrillo += round(float(ipo))

        elif 19 == tax and cat == "BOLSA":
            cls.tax_19_bolsa += round(float(value))

    @staticmethod
    def Tax_General(cls, number,date, client):
        if int(cls.tax_19_general) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 24080501,
                'accounting_account': "IVA VTA GNERAL 19%",
                'identification':client,
                'description':"IVA VTA GNERAL 19%",
                'debit':0,
                'credit':cls.tax_19_general
            })
        if int(cls.tax_5_general) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 24080502,
                'accounting_account': "IVA VTA GNERAL 5%",
                'identification':client,
                'description':"VA VTA GNERAL 5%",
                'debit':0,
                'credit':cls.tax_5_general
            })
        if int(cls.ipo_0_general) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 14300115,
                'accounting_account': "IMPOCO VTA GENERAL",
                'identification':client,
                'description':"IMPOCO VTA GENERAL",
                'debit':0,
                'credit':cls.ipo_0_general
            })

    @staticmethod
    def Licor(cls,number,date,client):
        if int(cls.tax_19_licor) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 24080505,
                'accounting_account': "IVA VTA LICOR 5",
                'identification':client,
                'description':"IVA VTA LICOR 5 ",
                'debit':0,
                'credit':cls.tax_19_licor
            })
        if int(cls.ipo_19_licor) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 14300116,
                'accounting_account': "IMPOCO VTA LICOR GNERAL",
                'identification':client,
                'description':"IMPOCO VTA LICOR GNERAL",
                'debit':0,
                'credit':cls.ipo_19_licor
            })

    @staticmethod    
    def Cerveza(cls,number,date,client):
        if int(cls.tax_19_cerveza) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 24080507,
                'accounting_account': "IVA VTA CERVEZA",
                'identification':client,
                'description':"IVA VTA CERVEZA",
                'debit':0,
                'credit':cls.tax_19_cerveza
            })
        if int(cls.ipo_19_cerveza) > 0:
            cls.data.append({
                'Secuencia':number,
                'Fecha elaboración': date,
                'Código contable': 14300117,
                'Cuenta contable': "IMPOCON VTA CERVEZA",
                'Identificación':client,
                'Descripción':"IMPOCON VTA CERVEZA",
                'Débito':0,
                'Crédito':cls.ipo_19_cerveza
            })

    @staticmethod
    def Vino(cls,number,date,client):
        if int(cls.tax_19_vino) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 24080509,
                'accounting_account': "IVA VTA VINO",
                'identification':client,
                'description':"IVA VTA VINO",
                'debit':0,
                'credit':cls.tax_19_vino
            })
        if int(cls.ipo_19_vino) > 0:
            cls.data.append({
                'Secuencia':number,
                'Fecha elaboración': date,
                'Código contable': 14300118,
                'Cuenta contable': "IMPOCON VTA VINO",
                'Identificación':client,
                'Descripción':"IMPOCON VTA VINO",
                'Débito':0,
                'Crédito':cls.ipo_19_vino
            })

    @staticmethod
    def Cigarrillo(cls,number,date,client):
        if int(cls.tax_19_cigarrillo) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 24080511,
                'accounting_account': "IVA 19 VTA CIGARRILLO",
                'identification':client,
                'description':"IVA 19 VTA CIGARRILLO",
                'debit':0,
                'credit':cls.tax_19_cigarrillo
            })
        if int(cls.ipo_19_cigarrillo) > 0:
            cls.data.append({
                'Secuencia':number,
                'Fecha elaboración': date,
                'Código contable': 14300119,
                'Cuenta contable': "IMPO VTA CIGARRILLO",
                'Identificación':client,
                'Descripción':"IMPO VTA CIGARRILLO",
                'Débito':0,
                'Crédito':cls.ipo_19_cigarrillo
            })

    @staticmethod
    def Gaseosa(cls,number,date,client):
        if int(cls.tax_19_gaseosa) > 0:
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 24080513,
                'accounting_account': "IVA GASEOSA",
                'identification':client,
                'description':"IVA GASEOSA",
                'debit':0,
                'credit':cls.tax_19_gaseosa
            })


    @staticmethod
    def segment_months(start, end):
        current_date = start
        end_date = end
        segments = []
        
        while current_date <= end_date:
            next_month = current_date.replace(day=28) + timedelta(days=4)
            last_day = next_month - timedelta(days=next_month.day)

            if last_day > end_date:
                last_day = end_date
            segments.append((current_date, last_day))
            current_date = last_day + timedelta(days=1)
        return segments

   
    @classmethod
    def generate_accounting(cls, data):
        result = False
        message = []
        # try:
        invoices_queryset = Invoice.objects.prefetch_related('details_invoice_set').filter(
            pk=data['pk_invoice'],
            branch=data['pk_branch']
        )
        cls.Taxes(cls)
        for invoice in invoices_queryset:
            total = 0
            number = invoice.number
            date = invoice.date
            client = invoice.customer.identification_number

            for detail in invoice.details_invoice_set.all():
                product = Product.objects.get(code=detail.code, branch=invoice.branch)
                cat = product.subcategory.category.name
                tax = detail.tax_value
                quantity = detail.quantity
                base = detail.price * quantity
                ipo = detail.ipo * quantity
                value_tax = detail.tax
                cls.Name_Category(cls, cat, tax, number, client, base, date)
                cls.Calculate_Taxes(cls,ipo,tax,value_tax,cat)
                cls.Tax_General(cls,number,date,client)
                cls.Gaseosa(cls,number,date,client)
                cls.Licor(cls,number,date,client)
                cls.Cerveza(cls,number,date,client)
                cls.Vino(cls,number,date,client)
                cls.Cigarrillo(cls,number,date,client)
                cls.Taxes(cls)
                total += base + value_tax + ipo
            cls.data.append({
                'sequence':number,
                'date_created': date,
                'accounting_code': 11050505,
                'accounting_account': f"Caja General {invoice.branch.name}",
                'identification':client,
                'description': f"Caja General {invoice.branch.name}",
                'debit': total,
                'credit':0
            })
            _result = cls.create_accountingentry(cls.data, invoice)
            if _result['result']:
                cls.data = []
                result = True
        # except Exception as e:
        #     message = str(e)
        #     print(e)
        return {'result': result, 'message': message, 'data':cls.data}


        
    @classmethod
    def generate_excel(cls, data):
        result = False
        message = None
        date_from_str = str(data.get('date_from'))
        date_to_str = str(data.get('date_to'))

        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d')
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d')
            if date_to < date_from:
                raise ValueError("The date 'date_to' cannot be earlier than 'date_from'.")
        except ValueError as e:
            raise ValueError("Dates must be in the format 'YYYY-MM-DD' and be valid.") from e

        # Obtiene los registros contables según el rango de fechas
        invoices_queryset = cls.objects.filter(
            invoice__date__range=(date_from_str, date_to_str),
            invoice__branch=data['pk_branch']
        )

        # Crea un libro de trabajo de Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Asiento Contable"

        # Escribir encabezados
        headers = ['Secuencia', 'Fecha elaboración', 'Código contable', 
                   'Cuenta contable', 'Identificación', 'Descripción', 
                   'Débito', 'Crédito']
        ws.append(headers)

        # Escribir datos
        for entry in invoices_queryset:
            ws.append([
                entry.sequence,
                entry.date_created,
                entry.accounting_code,
                entry.accounting_account,
                entry.identification,
                entry.description,
                entry.debit,
                entry.credit,
            ])

        # Guardar el archivo Excel
        documentation_dir = "C:/Users/Public/Videos/Nueva carpeta (2)/api/documentation"
        file_path = os.path.join(documentation_dir, 'asiento_contable.xlsx')
        wb.save(file_path)

        # Actualiza el resultado y mensaje
        result = True
        message = "El archivo Excel se ha creado y guardado exitosamente."

        return {'result': result, 'message': message}















































 # @classmethod
    # def generate_accounting(cls, data):
    #     result = False
    #     message = []
    #     date_from_str = str(data.get('date_from'))
    #     date_to_str = str(data.get('date_to'))
    #     try:
    #         date_from = datetime.strptime(date_from_str, '%Y-%m-%d')
    #         date_to = datetime.strptime(date_to_str, '%Y-%m-%d')
    #         if date_to < date_from:
    #             raise ValueError("The date 'date_to' cannot be earlier than 'date_from'.")
    #     except ValueError as e:
    #         raise ValueError("Dates must be in the format 'YYYY-MM-DD' and be valid.") from e
        
    #     months = cls.segment_months(date_from, date_to)
    #     invoices_queryset = Invoice.objects.prefetch_related('details_invoice_set').filter(
    #         date__range=(date_from_str, date_to_str),
    #         branch=data['pk_branch']
    #     )
    #     for i, (start, end) in enumerate(months, 1):
    #         monthly_invoices = invoices_queryset.filter(
    #             date__range=(start.strftime('%Y-%m-%d'), end.strftime('%Y-%m-%d'))
    #         )
    #         if monthly_invoices.exists():
    #             message.append(f"Invoices found for month {i} between {start} and {end}.")
    #             result = True
    #             for invoice in monthly_invoices:
    #                 for detail in invoice.details_invoice_set.all():
    #                     product = Product.objects.get(code=detail.code, branch=invoice.branch)
    #                     print(product)
    #                     cat = product.subcategory.category.name
    #                     tax = detail.tax_value
    #                     number = invoice.number
    #                     client = invoice.customer.name
    #                     quantity = detail.quantity
    #                     base = detail.price * quantity
    #                     date = invoice.date
    #                     ipo = detail.ipo * quantity
    #                     value_tax = detail.tax
    #                     cls.Name_Category(cls, cat, tax, number, client, base, date)
    #                     cls.Calculate_Taxes(cls,ipo,tax,value,cat)
    #                     print(data)
    #         else:
    #             message.append(f"No invoices found for month {i} between {start} and {end}.")
    #     return {'result': result, 'message': message}
